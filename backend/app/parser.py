import re
import shutil
import subprocess
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

import pymupdf
import pytesseract
from openpyxl import load_workbook
from PIL import Image

from .db import uid

ALLOWED = {".pdf", ".xlsx", ".xls", ".pptx", ".ppt", ".html", ".htm", ".png", ".jpg", ".jpeg", ".webp", ".txt"}
MAX_CHUNK_CHARS = 12_000


class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "svg", "nav", "footer", "noscript"}:
            self.hidden += 1
        elif not self.hidden and tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "article", "section"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "svg", "nav", "footer", "noscript"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip() + " ")

    def text(self):
        return re.sub(r"\n{3,}", "\n\n", "".join(self.parts)).strip()


def convert_legacy(path, suffix):
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise ValueError("旧形式の変換にLibreOfficeが必要です。XLSXまたはPPTXで保存して再登録してください")
    with tempfile.TemporaryDirectory() as temp:
        profile = Path(temp) / "profile"
        command = [
            executable,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx" if suffix == ".xls" else "pptx",
            "--outdir",
            temp,
            str(path),
        ]
        subprocess.run(command, check=True, capture_output=True, timeout=180)
        converted = Path(temp) / (path.stem + (".xlsx" if suffix == ".xls" else ".pptx"))
        if not converted.exists():
            raise ValueError("旧形式の変換に失敗しました。XLSXまたはPPTXで保存して再登録してください")
        return parse(converted)


def split_text(text, max_chars=MAX_CHUNK_CHARS):
    """Split extracted text on readable boundaries without dropping content."""
    text = text.strip()
    if not text:
        return []
    parts = []
    start = 0
    while start < len(text):
        hard_end = min(start + max_chars, len(text))
        end = hard_end
        if hard_end < len(text):
            window = text[start:hard_end]
            minimum = max_chars // 2
            candidates = [
                window.rfind("\n\n", minimum),
                window.rfind("\n", minimum),
                window.rfind("。", minimum),
                window.rfind(". ", minimum),
                window.rfind(" ", minimum),
            ]
            boundary = max(candidates)
            if boundary >= minimum:
                marker = window[boundary : boundary + 2]
                end = start + boundary + (2 if marker in ("\n\n", ". ") else 1)
        piece = text[start:end].strip()
        if piece:
            parts.append(piece)
        start = end
    return parts


def split_text_with_lines(text, max_chars=MAX_CHUNK_CHARS, first_line=1):
    """Split text while keeping inclusive source line numbers for each piece."""
    raw = str(text)
    pieces = split_text(raw, max_chars)
    result = []
    cursor = 0
    for piece in pieces:
        start = raw.find(piece, cursor)
        if start < 0:
            start = cursor
        end = start + len(piece)
        line_start = first_line + raw[:start].count("\n")
        line_end = first_line + raw[:end].count("\n")
        char_start = char_end = None
        if line_start == line_end:
            line_offset = raw.rfind("\n", 0, start) + 1
            char_start = start - line_offset + 1
            char_end = end - line_offset
        result.append(
            (piece, line_start, max(line_start, line_end), char_start, char_end)
        )
        cursor = end
    return result


def split_chunks(chunks, max_chars=MAX_CHUNK_CHARS):
    """Bound every extracted chunk while retaining its page/sheet provenance."""
    result = []
    for chunk in chunks:
        original_id = chunk.get("id") or uid()
        text = str(chunk.get("text", ""))
        first_line = chunk.get("line_start")
        if first_line is None:
            cell_range = str(chunk.get("cell_range") or "")
            match = re.search(r"[A-Z]+(\d+)", cell_range)
            first_line = int(match.group(1)) if match else 1
        pieces = split_text_with_lines(text, max_chars, int(first_line))
        base_char = chunk.get("char_start")
        if base_char is not None:
            pieces = [
                (
                    piece,
                    line_start,
                    line_end,
                    char_start + int(base_char) - 1 if char_start is not None else None,
                    char_end + int(base_char) - 1 if char_end is not None else None,
                )
                for piece, line_start, line_end, char_start, char_end in pieces
            ]
        if len(pieces) <= 1:
            if pieces:
                piece, line_start, line_end, char_start, char_end = pieces[0]
                result.append(
                    {
                        **chunk,
                        "id": original_id,
                        "text": piece,
                        "line_start": line_start,
                        "line_end": line_end,
                        "char_start": char_start,
                        "char_end": char_end,
                    }
                )
            continue
        for index, (piece, line_start, line_end, char_start, char_end) in enumerate(pieces, start=1):
            result.append(
                {
                    **chunk,
                    "id": f"{original_id}:segment:{index}",
                    "text": piece,
                    "line_start": line_start,
                    "line_end": line_end,
                    "char_start": char_start,
                    "char_end": char_end,
                    "source_chunk_id": original_id,
                    "segment_index": index,
                    "segment_count": len(pieces),
                }
            )
    return result


def ocr(image):
    result = pytesseract.image_to_data(
        image, lang="jpn+eng", output_type=pytesseract.Output.DICT, timeout=90
    )
    words = []
    previous_line = None
    for text, confidence, block, paragraph, line in zip(
        result["text"],
        result["conf"],
        result["block_num"],
        result["par_num"],
        result["line_num"],
    ):
        if not text.strip() or float(confidence) < 0:
            continue
        line_key = (block, paragraph, line)
        if previous_line is not None:
            words.append("\n" if line_key != previous_line else " ")
        words.append(text)
        previous_line = line_key
    confidences = [
        float(confidence)
        for text, confidence in zip(result["text"], result["conf"])
        if text.strip() and float(confidence) >= 0
    ]
    return "".join(words), round(
        sum(confidences) / len(confidences), 1
    ) if confidences else 0


def parse(path):
    suffix = path.suffix.lower()
    if suffix in {".xls", ".ppt"}:
        return convert_legacy(path, suffix)
    chunks = []

    def add(text, **loc):
        if text.strip():
            raw_text = str(text)
            cleaned = raw_text.strip()
            leading_whitespace = raw_text[: len(raw_text) - len(raw_text.lstrip())]
            line_start = int(loc.get("line_start", 1)) + leading_whitespace.count("\n")
            chunks.append(
                {
                    "id": uid(),
                    "text": cleaned,
                    "categories": [],
                    **loc,
                    "line_start": line_start,
                    "line_end": line_start + max(0, len(cleaned.splitlines()) - 1),
                }
            )

    if suffix == ".txt":
        add(path.read_text(encoding="utf-8-sig"))
    elif suffix in {".html", ".htm"}:
        html = VisibleHTML()
        raw = path.read_bytes()
        head = raw[:4096].decode("ascii", errors="ignore")
        charset = re.search(r"charset\s*=\s*['\"]?([\w-]+)", head, re.IGNORECASE)
        encoding = charset.group(1) if charset else "utf-8"
        try:
            html.feed(raw.decode(encoding, errors="replace"))
        except LookupError:
            html.feed(raw.decode("utf-8", errors="replace"))
        add(html.text())
    elif suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        for sheet in wb:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            max_row = sheet.max_row or 0
            max_column = sheet.max_column or 0
            if max_row > 250000 or max_column > 2000:
                raise ValueError("シートが大きすぎます（上限250,000行・2,000列）。分けて登録してください")
            for start in range(1, max_row + 1, 40):
                end = min(start + 39, max_row)
                rows = [
                    "\t".join(
                        "" if v is None else re.sub(r"\r\n|\r|\n", " ", str(v))
                        for v in row
                    )
                    for row in sheet.iter_rows(
                        min_row=start, max_row=end, values_only=True
                    )
                ]
                from openpyxl.utils import get_column_letter

                add(
                    "\n".join(rows),
                    sheet_name=sheet.title,
                    cell_range=f"A{start}:{get_column_letter(max_column)}{end}",
                    line_start=start,
                    line_end=end,
                    table=True,
                )
        wb.close()
    elif suffix == ".pptx":
        with zipfile.ZipFile(path) as presentation:
            slides = [
                name for name in presentation.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ]
            slides.sort(key=lambda name: int(re.search(r"slide(\d+)", name).group(1)))
            if not slides:
                raise ValueError("PowerPointのスライドを読み取れませんでした")
            for name in slides:
                info = presentation.getinfo(name)
                if info.file_size > 20 * 1024 * 1024:
                    raise ValueError("1スライドの展開サイズが大きすぎます")
                root = ElementTree.fromstring(presentation.read(name))
                lines = [
                    node.text.strip() for node in root.iter()
                    if node.tag.endswith("}t") and node.text and node.text.strip()
                ]
                add("\n".join(lines), slide_number=int(re.search(r"slide(\d+)", name).group(1)))
    elif suffix == ".pdf":
        with pymupdf.open(path) as document:
            for index, page in enumerate(document):
                text = page.get_text("text", sort=True) or ""
                if len(text.strip()) > 20:
                    add(text, page_number=index + 1)
                else:
                    try:
                        with tempfile.TemporaryDirectory() as temp:
                            out = Path(temp) / "page"
                            subprocess.run(
                                [
                                    "pdftoppm", "-f", str(index + 1), "-l", str(index + 1),
                                    "-scale-to", "2400", "-singlefile", "-png", str(path), str(out),
                                ],
                                check=True, capture_output=True, timeout=90,
                            )
                            with Image.open(str(out) + ".png") as im:
                                ocr_text, confidence = ocr(im)
                        add(ocr_text or text, page_number=index + 1, ocr_confidence=confidence)
                    except (OSError, subprocess.SubprocessError, RuntimeError):
                        add(text, page_number=index + 1)
    else:
        with Image.open(path) as im:
            text, confidence = ocr(im)
        add(text, page_number=1, ocr_confidence=confidence)
    if not chunks:
        raise ValueError(
            "文字を抽出できませんでした。鮮明な資料を使用するかテキストで登録してください"
        )
    return chunks
